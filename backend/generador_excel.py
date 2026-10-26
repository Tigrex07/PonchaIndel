from datetime import datetime, timedelta, time
import os
import subprocess
import sys
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.modelos import obtener_horario_maestro_por_fecha

class GeneradorExcelProgramatico:

  def __init__(self):
    pass

  def generar_control_quincenal(
      self,
      departamento,
      periodo_texto,
      fecha_inicio_str,
      fecha_fin_str,
      df_registros,
      ruta_salida,
  ):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control Quincenal"
    ws.views.sheetView[0].showGridLines = True

    font_titulo = Font(name="Arial", size=13, bold=True)
    font_sub = Font(name="Arial", size=10, bold=True)
    font_header = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    font_contenido = Font(name="Arial", size=9)
    font_alerta = Font(name="Arial", size=8, bold=True, color="990000")
    font_notas = Font(name="Arial", size=8, italic=True)
    font_simbologia = Font(name="Arial", size=8, bold=True)

    fill_negro = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    fill_gris = PatternFill(start_color="374151", end_color="374151", fill_type="solid")

    borde_fino = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ruta_logo_utn = "assets/logos/logo-utn.png"
    ruta_logo_iqs = "assets/logos/logo-iqs.png"

    if os.path.exists(ruta_logo_utn):
      img_utn = Image(ruta_logo_utn)
      img_utn.width = 100
      img_utn.height = 42
      ws.add_image(img_utn, "A1")

    if os.path.exists(ruta_logo_iqs):
      img_iqs = Image(ruta_logo_iqs)
      img_iqs.width = 45
      img_iqs.height = 42

    ws.merge_cells("C1:V1")
    ws["C1"] = "Control Quincenal de Asistencia Administrativo y PTC"
    ws["C1"].font = font_titulo
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C3:V3")
    ws["C3"] = f"Area/Dirección: {departamento}"
    ws["C3"].font = font_sub
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C4:V4")
    ws["C4"] = f"Quincena del: {periodo_texto}"
    ws["C4"].font = Font(name="Arial", size=9, bold=True, color="2563EB")
    ws["C4"].alignment = Alignment(horizontal="center", vertical="center")

    dt_inicio = datetime.strptime(fecha_inicio_str, "%Y-%m-%d")
    dt_fin = datetime.strptime(fecha_fin_str, "%Y-%m-%d")

    dias_quincena = []
    actual = dt_inicio
    while actual <= dt_fin:
      dias_quincena.append(actual)
      actual += timedelta(days=1)

    ws.cell(row=6, column=1, value="NOMBRE").font = font_header
    ws.cell(row=6, column=1).fill = fill_negro
    ws.cell(row=6, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A6:A8")

    for idx, dt_dia in enumerate(dias_quincena):
      col_e = 2 + (idx * 2)
      col_s = 3 + (idx * 2)
      letra_e = get_column_letter(col_e)
      letra_s = get_column_letter(col_s)

      ws.merge_cells(f"{letra_e}6:{letra_s}6")
      cell_dia = ws.cell(row=6, column=col_e, value=dt_dia.day)
      cell_dia.font = font_header
      cell_dia.fill = fill_negro
      cell_dia.alignment = Alignment(horizontal="center", vertical="center")

      ws.cell(row=7, column=col_e, value="Entrada").font = font_header
      ws.cell(row=7, column=col_e).fill = fill_gris
      ws.cell(row=7, column=col_e).alignment = Alignment(horizontal="center", vertical="center")

      ws.cell(row=7, column=col_s, value="Salida").font = font_header
      ws.cell(row=7, column=col_s).fill = fill_gris
      ws.cell(row=7, column=col_s).alignment = Alignment(horizontal="center", vertical="center")

    col_total_idx = 2 + (len(dias_quincena) * 2)
    col_total_letra = get_column_letter(col_total_idx)
    ws.cell(row=6, column=col_total_idx, value="TOTAL HRS\nEFECTIVAS").font = font_header
    ws.cell(row=6, column=col_total_idx).fill = fill_negro
    ws.cell(row=6, column=col_total_idx).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(f"{col_total_letra}6:{col_total_letra}8")

    if os.path.exists(ruta_logo_iqs):
      col_logo_letra = get_column_letter(max(2, col_total_idx - 1))
      ws.add_image(img_iqs, f"{col_logo_letra}1")

    maestros_unicos = []
    if df_registros is not None and not df_registros.empty:
      if "Nombre" in df_registros.columns:
        maestros_unicos = df_registros["Nombre"].unique().tolist()

    fila_inicio = 9
    for idx_m, nombre_docente in enumerate(maestros_unicos):
      f_actual = fila_inicio + idx_m
      ws.row_dimensions[f_actual].height = 28

      cell_nom = ws.cell(row=f_actual, column=1, value=nombre_docente)
      cell_nom.font = font_contenido
      cell_nom.border = borde_fino
      cell_nom.alignment = Alignment(horizontal="left", vertical="center")

      df_docente = df_registros[df_registros["Nombre"] == nombre_docente]
      minutos_totales_trabajados = 0

      for idx_d, dt_dia in enumerate(dias_quincena):
        col_e = 2 + (idx_d * 2)
        col_s = 3 + (idx_d * 2)

        cell_e = ws.cell(row=f_actual, column=col_e)
        cell_s = ws.cell(row=f_actual, column=col_s)

        cell_e.border = borde_fino
        cell_e.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_e.font = font_contenido

        cell_s.border = borde_fino
        cell_s.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_s.font = font_contenido

        # --- OBTENER HORARIO DINÁMICO POR FECHA Y CUATRIMESTRE ---
        fecha_dia_str = dt_dia.strftime("%Y-%m-%d")
        horario_asignado = obtener_horario_maestro_por_fecha(nombre_docente, fecha_dia_str)
        
        if horario_asignado and horario_asignado[0] and horario_asignado[1]:
            hora_oficial_str = str(horario_asignado[0]).strip()
            hora_oficial_salida_str = str(horario_asignado[1]).strip()
        else:
            hora_oficial_str = "08:00:00"
            hora_oficial_salida_str = "16:00:00"

        if len(hora_oficial_str.split(':')) == 2:
            hora_oficial_str += ":00"
        if len(hora_oficial_salida_str.split(':')) == 2:
            hora_oficial_salida_str += ":00"

        t_oficial = datetime.strptime(hora_oficial_str, "%H:%M:%S").time()
        dt_oficial = datetime.combine(datetime.today(), t_oficial)
        dt_tolerancia = dt_oficial + timedelta(minutes=10)
        t_tolerancia = dt_tolerancia.time()
        
        t_oficial_sal = datetime.strptime(hora_oficial_salida_str, "%H:%M:%S").time()
        # --------------------------------------------------------

        sub_dia = df_docente[df_docente["fecha"].apply(lambda x: self._es_mismo_dia(x, dt_dia))]
        entradas = sub_dia[sub_dia["tipo"] == "ENTRADA"]["hora"].values
        salidas = sub_dia[sub_dia["tipo"] == "SALIDA"]["hora"].values

        es_fin_de_semana = dt_dia.weekday() >= 5
        t_ing = None
        t_sal = None

        if len(entradas) > 0:
          h_entrada_str = str(entradas[0]).strip()
          
          if len(h_entrada_str.split(':')) == 2:
              h_entrada_str += ":00"
              
          try:
            t_ing = datetime.strptime(h_entrada_str, "%H:%M:%S").time()
            if t_ing > t_tolerancia:
              dt_ing_val = datetime.combine(datetime.today(), t_ing)
              min_ret_dia = int((dt_ing_val - dt_tolerancia).total_seconds() / 60)
              cell_e.value = f"{h_entrada_str}\n(+{min_ret_dia}m Ret)"
              cell_e.font = font_alerta
            else:
              cell_e.value = h_entrada_str
          except ValueError:
            cell_e.value = h_entrada_str
        else:
          if es_fin_de_semana:
            cell_e.value = ""
          else:
            cell_e.value = "FIE"
            cell_e.font = font_alerta

        if len(salidas) > 0:
          h_salida_str = str(salidas[0]).strip()
          
          if len(h_salida_str.split(':')) == 2:
              h_salida_str += ":00"
              
          try:
            t_sal = datetime.strptime(h_salida_str, "%H:%M:%S").time()
            if t_sal < t_oficial_sal:
              cell_s.value = f"{h_salida_str}\n(ST)"
              cell_s.font = font_alerta
            else:
              cell_s.value = h_salida_str
          except ValueError:
            cell_s.value = h_salida_str
        else:
          if len(entradas) > 0 and not es_fin_de_semana:
            cell_s.value = "FIS"
            cell_s.font = font_alerta
          else:
            cell_s.value = ""

        if t_ing and t_sal:
            dt_i = datetime.combine(datetime.today(), t_ing)
            dt_s = datetime.combine(datetime.today(), t_sal)
            if dt_s < dt_i:
                dt_s += timedelta(days=1)
            diff_secs = (dt_s - dt_i).total_seconds()
            if diff_secs > 0:
                minutos_totales_trabajados += int(diff_secs / 60)

      hrs_tot = minutos_totales_trabajados // 60
      mins_tot = minutos_totales_trabajados % 60
      cell_tot = ws.cell(row=f_actual, column=col_total_idx, value=f"{hrs_tot} hrs\n{mins_tot} min")
      cell_tot.border = borde_fino
      cell_tot.font = Font(name="Arial", size=9, bold=True)
      cell_tot.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    f_notas = fila_inicio + max(len(maestros_unicos), 1) + 2
    ws.cell(row=f_notas, column=1, value="En este Control Quincenal aparecen las entradas registradas después del horario estipulado.").font = font_notas
    ws.cell(row=f_notas + 1, column=1, value="Se consideran 10 minutos de TOLERANCIA solo en el registro de las Entradas, según el Contrato Colectivo.").font = font_notas
    ws.cell(row=f_notas + 2, column=1, value="Después de los 10 minutos de TOLERANCIA se comenzarán a contar los minutos de Retardo del Personal.").font = font_notas
    ws.cell(row=f_notas + 3, column=1, value="ST = Salida Temprana (Registrada antes del horario estipulado de salida).").font = font_notas

    f_simb = f_notas + 6
    simbologia_texto = (
        "FI = Falta Injustificada    I = Incapacidad    AM = Atención Médica    P = Permiso    C = Comisión\n"
        "FIS = Falta, Salida No Registrada    FIE = Falta, Entrada No Registrada    PS = Permiso Sin Goce    ST = Salida Temprana"
    )
    ws.cell(row=f_simb, column=1, value=simbologia_texto).font = font_simbologia
    ws.cell(row=f_simb + 4, column=1, value="F01PGRRH02.01").font = Font(name="Arial", size=9, bold=True)

    ws.column_dimensions["A"].width = 35
    wb.save(ruta_salida)
    return ruta_salida

  def _es_mismo_dia(self, fecha_db, dt_ref):
    if not fecha_db: return False
    try:
      f_str = str(fecha_db).strip().split(" ")[0]
      dt = datetime.strptime(f_str, "%Y-%m-%d")
      return dt.date() == dt_ref.date()
    except ValueError:
      return False

  def abrir_documento_generado(self, ruta_archivo):
    if sys.platform == "win32":
      os.startfile(ruta_archivo)
    elif sys.platform == "darwin":
      subprocess.call(["open", ruta_archivo])
    else:
      subprocess.call(["xdg-open", ruta_archivo])