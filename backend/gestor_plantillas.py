# backend/gestor_plantillas.py
from datetime import datetime
import os
import subprocess
import sys
import openpyxl


class GestorPlantillas:

  def __init__(self, ruta_assets="assets/plantillas"):
    self.ruta_assets = ruta_assets

  def _obtener_ruta_plantilla(self, nombre_archivo):
    ruta_directa = os.path.join(self.ruta_assets, nombre_archivo)
    if os.path.exists(ruta_directa):
      return ruta_directa
    ruta_relativa = os.path.join("..", self.ruta_assets, nombre_archivo)
    if os.path.exists(ruta_relativa):
      return ruta_relativa
    return ruta_directa

  def _escribir_seguro(self, hoja, coordenada, valor):
    """Escribe un valor asegurando no afectar celdas combinadas generales."""
    try:
      celda = hoja[coordenada]
      if type(celda).__name__ == "MergedCell":
        for rango in hoja.merged_cells.ranges:
          if coordenada in rango:
            hoja.cell(row=rango.min_row, column=rango.min_col, value=valor)
            return
      else:
        celda.value = valor
    except Exception:
      pass

  def inyectar_control_quincenal(
      self, departamento, periodo, df_registros, ruta_salida
  ):
    """Inyecta los datos de los maestros, marca retardos y asigna 'FI' (Falta Injustificada) si no hay asistencia."""
    nombre_plantilla = (
        "F01PGRRH02.01 Formato Control Quincenal de Asistencia Administrativos"
        "  y PTC.xlsx"
    )
    ruta_base = self._obtener_ruta_plantilla(nombre_plantilla)

    if not os.path.exists(ruta_base):
      raise FileNotFoundError(
          f"No se encontró la plantilla oficial en: {ruta_base}"
      )

    wb = openpyxl.load_workbook(ruta_base)
    ws = wb.active

    # 1. Inyectar Encabezados Generales
    self._escribir_seguro(ws, "A3", departamento)
    self._escribir_seguro(ws, "A5", periodo)

    # 2. Configuración de Filas y Mapeo de Fechas
    fila_inicio = 11

    maestros_unicos = []
    fechas_globales_objs = []

    if df_registros is not None and not df_registros.empty:
      if "Nombre" in df_registros.columns:
        maestros_unicos = df_registros["Nombre"].unique().tolist()

      if "fecha" in df_registros.columns:
        fechas_crudas = df_registros["fecha"].dropna().unique()
        lista_dt = []
        for f in fechas_crudas:
          f_str = str(f).strip().split(" ")[0]
          for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
              lista_dt.append(datetime.strptime(f_str, fmt))
              break
            except ValueError:
              continue

        fechas_globales_objs = sorted(list(set(lista_dt)))[:15]

        # Pintar los días en la Fila 8
        for d_idx, dt_val in enumerate(fechas_globales_objs):
          col_entrada = 2 + (d_idx * 2)
          if col_entrada < ws.max_column:
            ws.cell(row=8, column=col_entrada, value=dt_val.day)

    total_maestros = len(maestros_unicos)

    if total_maestros > 0:
      if total_maestros > 1:
        ws.insert_rows(fila_inicio + 1, amount=total_maestros - 1)

      for idx, nombre_docente in enumerate(maestros_unicos):
        f_actual = fila_inicio + idx
        ws.row_dimensions[f_actual].height = 25

        # Copiar formato de la fila base
        if idx > 0:
          for col in range(1, ws.max_column + 1):
            cell_src = ws.cell(row=fila_inicio, column=col)
            cell_dst = ws.cell(row=f_actual, column=col)
            cell_dst.font = (
                cell_src.font.copy() if cell_src.font else None
            )
            cell_dst.fill = (
                cell_src.fill.copy() if cell_src.fill else None
            )
            cell_dst.border = (
                cell_src.border.copy() if cell_src.border else None
            )
            cell_dst.alignment = (
                cell_src.alignment.copy() if cell_src.alignment else None
            )
            cell_dst.number_format = cell_src.number_format

        # Escribir el Nombre del Docente
        self._escribir_seguro(ws, f"A{f_actual}", nombre_docente)

        df_docente = df_registros[df_registros["Nombre"] == nombre_docente]

        if fechas_globales_objs:
          for d_idx, dt_global in enumerate(fechas_globales_objs):
            sub_fecha = df_docente[
                df_docente["fecha"].apply(
                    lambda x: self._coincide_dia(x, dt_global)
                )
            ]

            entradas = sub_fecha[sub_fecha["tipo"] == "ENTRADA"][
                "hora"
            ].values
            salidas = sub_fecha[sub_fecha["tipo"] == "SALIDA"]["hora"].values

            col_entrada = 2 + (d_idx * 2)
            col_salida = 3 + (d_idx * 2)

            # EVALUACIÓN DE INCIDENCIAS:
            if len(entradas) > 0:
              hora_e = str(entradas[0])
              # Detectar retardo si la entrada es mayor a las 08:10:00 (ejemplo basado en la tolerancia oficial)
              if hora_e > "08:10:00":
                hora_e += " (Retardo)"
              if col_entrada < ws.max_column:
                ws.cell(row=f_actual, column=col_entrada, value=hora_e)
            else:
              # Si no hay registro de entrada, marcar Falta Injustificada (FI)
              if col_entrada < ws.max_column:
                ws.cell(row=f_actual, column=col_entrada, value="FI")

            if len(salidas) > 0:
              hora_s = str(salidas[0])
              if col_salida < ws.max_column:
                ws.cell(row=f_actual, column=col_salida, value=hora_s)
            else:
              if col_salida < ws.max_column:
                ws.cell(row=f_actual, column=col_salida, value="")

    wb.save(ruta_salida)
    return ruta_salida

  def _coincide_dia(self, fecha_db, dt_referencia):
    if not fecha_db:
      return False
    f_str = str(fecha_db).strip().split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
      try:
        dt = datetime.strptime(f_str, fmt)
        if (
            dt.year == dt_referencia.year
            and dt.month == dt_referencia.month
            and dt.day == dt_referencia.day
        ):
          return True
        break
      except ValueError:
        continue
    return False

  def abrir_documento_generado(self, ruta_archivo):
    if sys.platform == "win32":
      os.startfile(ruta_archivo)
    elif sys.platform == "darwin":
      subprocess.call(["open", ruta_archivo])
    else:
      subprocess.call(["xdg-open", ruta_archivo])