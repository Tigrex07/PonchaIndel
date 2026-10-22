import os
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
import pandas as pd

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.modelos import guardar_ponchadas_desde_df
from frontend.tema import COLORES

class VistaPonchadas(ctk.CTkFrame):
    def __init__(self, parent, on_procesar_exito=None):
        super().__init__(parent, fg_color=COLORES["fondo_principal"])
        self.on_procesar_exito = on_procesar_exito
        self.ruta_archivo = None
        self.crear_interfaz()

    def crear_interfaz(self):
        header_frame = ctk.CTkFrame(self, fg_color="transparent")
        header_frame.pack(fill="x", padx=20, pady=(15, 5))

        ctk.CTkLabel(
            header_frame, text="✨ Sincronización de Ponchadora", 
            font=ctk.CTkFont(size=26, weight="bold"), text_color=COLORES["texto_principal"]
        ).pack(side="left")

        self.status_bd = ctk.CTkLabel(
            header_frame, text="🟢 Conectado", font=ctk.CTkFont(size=14, weight="bold"),
            text_color=COLORES["exito"], fg_color="#064E3B", corner_radius=10, padx=12, pady=6
        )
        self.status_bd.pack(side="right")

        main_frame = ctk.CTkFrame(self, fg_color="transparent")
        main_frame.pack(fill="both", expand=True, padx=20, pady=10)

        # PANEL IZQUIERDO (Con borde resaltado)
        left_frame = ctk.CTkFrame(
            main_frame, width=400, fg_color=COLORES["panel_secundario"], 
            corner_radius=15, border_color=COLORES["borde_acento"], border_width=2
        )
        left_frame.pack(side="left", fill="both", padx=(0, 15), pady=5)

        ctk.CTkLabel(
            left_frame, text="📂 Cargar Archivo CSV", font=ctk.CTkFont(size=18, weight="bold"), text_color=COLORES["texto_principal"]
        ).pack(side="top", pady=(20, 5), padx=20, anchor="w")
        
        ctk.CTkLabel(
            left_frame, text="Seleccione el archivo exportado del checador.", font=ctk.CTkFont(size=13), text_color=COLORES["texto_secundario"]
        ).pack(side="top", anchor="w", padx=20, pady=(0, 15))

        file_box = ctk.CTkFrame(left_frame, fg_color=COLORES["panel_terciario"], corner_radius=10)
        file_box.pack(side="top", fill="x", padx=20, pady=5)

        self.entry_ruta = ctk.CTkEntry(
            file_box, placeholder_text="Ningún archivo seleccionado...", fg_color="transparent",
            text_color=COLORES["texto_principal"], border_width=0, height=45, font=ctk.CTkFont(size=13)
        )
        self.entry_ruta.pack(side="left", fill="x", expand=True, padx=10)

        self.btn_buscar = ctk.CTkButton(
            file_box, text="🔍 Buscar", width=90, height=35, corner_radius=8,
            fg_color=COLORES["primario"], hover_color=COLORES["primario_hover"],
            font=ctk.CTkFont(size=13, weight="bold"), command=self.seleccionar_archivo
        )
        self.btn_buscar.pack(side="right", padx=8, pady=8)

        # ANCLAMOS LOS BOTONES AL FONDO PARA QUE NO SE PIERDAN
        self.label_estado = ctk.CTkLabel(left_frame, text="Esperando archivo...", font=ctk.CTkFont(size=13), text_color=COLORES["texto_secundario"])
        self.label_estado.pack(side="bottom", pady=(0, 20))

        self.btn_procesar = ctk.CTkButton(
            left_frame, text="🚀 Procesar, Sincronizar y Exportar Excel", font=ctk.CTkFont(size=14, weight="bold"),
            fg_color=COLORES["exito"], hover_color=COLORES["exito_hover"], text_color="#FFFFFF",
            height=50, corner_radius=10, command=self.procesar_csv
        )
        self.btn_procesar.pack(side="bottom", fill="x", padx=20, pady=(10, 5))

        # PANEL DERECHO
        right_frame = ctk.CTkFrame(main_frame, fg_color=COLORES["panel_secundario"], corner_radius=15, border_color=COLORES["bordes"], border_width=1)
        right_frame.pack(side="right", fill="both", expand=True, pady=5)

        ctk.CTkLabel(
            right_frame, text="📈 Resumen de Sincronización", font=ctk.CTkFont(size=18, weight="bold"), text_color=COLORES["texto_principal"]
        ).pack(pady=(20, 15), padx=25, anchor="w")

        self.card_container = ctk.CTkFrame(right_frame, fg_color=COLORES["panel_terciario"], corner_radius=12)
        self.card_container.pack(fill="both", expand=True, padx=25, pady=(0, 25))

        top_box = ctk.CTkFrame(self.card_container, fg_color="transparent")
        top_box.pack(fill="x", padx=25, pady=(25, 15))

        self.lbl_icon = ctk.CTkLabel(
            top_box, text="📄", font=ctk.CTkFont(size=28), text_color="#FFFFFF", 
            fg_color=COLORES["panel_secundario"], width=60, height=60, corner_radius=30
        )
        self.lbl_icon.pack(side="left", padx=(0, 20))

        info_meta = ctk.CTkFrame(top_box, fg_color="transparent")
        info_meta.pack(side="left", fill="both", expand=True)

        self.lbl_nombre_archivo = ctk.CTkLabel(info_meta, text="Sin archivo cargado", font=ctk.CTkFont(size=18, weight="bold"), text_color=COLORES["texto_principal"], anchor="w")
        self.lbl_nombre_archivo.pack(fill="x")
        self.lbl_peso_archivo = ctk.CTkLabel(info_meta, text="Ruta: N/A", font=ctk.CTkFont(size=12), text_color=COLORES["texto_secundario"], anchor="w")
        self.lbl_peso_archivo.pack(fill="x")

        metrics_box = ctk.CTkFrame(self.card_container, fg_color=COLORES["fondo_principal"], corner_radius=10)
        metrics_box.pack(fill="x", padx=25, pady=20)

        m1 = ctk.CTkFrame(metrics_box, fg_color="transparent")
        m1.pack(side="left", fill="both", expand=True, padx=15, pady=15)
        ctk.CTkLabel(m1, text="REGISTROS", font=ctk.CTkFont(size=12, weight="bold"), text_color=COLORES["texto_secundario"]).pack()
        self.lbl_stat_filas = ctk.CTkLabel(m1, text="0", font=ctk.CTkFont(size=32, weight="bold"), text_color=COLORES["primario"])
        self.lbl_stat_filas.pack()

        m2 = ctk.CTkFrame(metrics_box, fg_color="transparent")
        m2.pack(side="right", fill="both", expand=True, padx=15, pady=15)
        ctk.CTkLabel(m2, text="DOCENTES", font=ctk.CTkFont(size=12, weight="bold"), text_color=COLORES["texto_secundario"]).pack()
        self.lbl_stat_docentes = ctk.CTkLabel(m2, text="0", font=ctk.CTkFont(size=32, weight="bold"), text_color=COLORES["primario"])
        self.lbl_stat_docentes.pack()

        self.lbl_status_footer = ctk.CTkLabel(self.card_container, text="🔴 Esperando acción del usuario.", font=ctk.CTkFont(size=14, weight="bold"), text_color=COLORES["peligro"])
        self.lbl_status_footer.pack(pady=(10, 20))

    def seleccionar_archivo(self):
        archivo = filedialog.askopenfilename(title="Selecciona el CSV", filetypes=[("Archivos CSV", "*.csv"), ("Todos", "*.*")])
        if archivo:
            self.ruta_archivo = archivo
            nombre = os.path.basename(archivo)
            self.entry_ruta.delete(0, tk.END)
            self.entry_ruta.insert(0, nombre)
            self.lbl_nombre_archivo.configure(text=nombre)
            self.lbl_peso_archivo.configure(text=f"Ruta: ...{archivo[-35:]}")
            self.label_estado.configure(text="¡Archivo listo para procesar!", text_color=COLORES["primario"])
            self.lbl_icon.configure(fg_color=COLORES["primario"])

    def procesar_csv(self):
        if not self.ruta_archivo:
            messagebox.showwarning("Atención", "Selecciona un archivo CSV.")
            return
        try:
            try:
                df = pd.read_csv(self.ruta_archivo, sep=";", encoding="utf-8-sig", on_bad_lines="skip")
            except:
                df = pd.read_csv(self.ruta_archivo, sep=";", encoding="latin1", on_bad_lines="skip")

            if df.shape[1] < 4:
                messagebox.showerror("Error", "Formato de archivo incorrecto.")
                return

            df_sub = df.iloc[:, [0, 3]].copy()
            df_sub.columns = ["Fecha_Hora_Raw", "Nombre"]
            df_sub["Nombre"] = df_sub["Nombre"].astype(str).str.strip()
            df_sub = df_sub.dropna(subset=["Fecha_Hora_Raw", "Nombre"])

            df_sub["Fecha_Hora"] = pd.to_datetime(df_sub["Fecha_Hora_Raw"].astype(str).str.strip(), format="%d/%m/%Y %H:%M:%S", errors="coerce")
            df_sub = df_sub.dropna(subset=["Fecha_Hora"])
            df_sub["Fecha_Texto"] = df_sub["Fecha_Hora"].dt.strftime("%Y-%m-%d")

            df_resumen = df_sub.groupby(["Fecha_Texto", "Nombre"], sort=False).agg(
                Entrada=("Fecha_Hora", lambda x: x.min().strftime("%H:%M:%S")),
                Salida=("Fecha_Hora", lambda x: x.max().strftime("%H:%M:%S"))
            ).reset_index()
            df_resumen.rename(columns={"Fecha_Texto": "Fecha"}, inplace=True)
            df_resumen = df_resumen[["Nombre", "Fecha", "Entrada", "Salida"]]

            self.lbl_stat_filas.configure(text=str(len(df_resumen)))
            self.lbl_stat_docentes.configure(text=str(df_resumen["Nombre"].nunique()))

            # Guardar en Base de Datos (Esto crea automáticamente los cuatrimestres y guarda maestros/ponchadas)
            guardar_ponchadas_desde_df(df_resumen)
            
            # --- EXCEL LIMPIO ---
            nombre_base = os.path.splitext(os.path.basename(self.ruta_archivo))[0]
            ruta_salida = filedialog.asksaveasfilename(
                title="Guardar Excel Formateado como...",
                initialfile=f"Reporte_Limpio_{nombre_base}.xlsx",
                defaultextension=".xlsx",
                filetypes=[("Archivos de Excel", "*.xlsx")],
            )

            if ruta_salida:
                with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
                    df_resumen.to_excel(writer, sheet_name="Asistencia", index=False, startrow=3)
                    ws = writer.sheets["Asistencia"]
                    ws.views.sheetView[0].showGridLines = True

                    font_banner = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
                    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                    font_data = Font(name="Calibri", size=11, color="1F2937")

                    fill_banner = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
                    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
                    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

                    border_thin = Side(border_style="thin", color="D9D9D9")
                    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

                    ws.merge_cells("A1:D2")
                    banner_cell = ws["A1"]
                    banner_cell.value = "REPORTE CONSOLIDADO DE ASISTENCIA Y PONCHADAS"
                    banner_cell.font = font_banner
                    banner_cell.fill = fill_banner
                    banner_cell.alignment = align_center

                    for col_idx in range(1, 5):
                        c = ws.cell(row=4, column=col_idx)
                        c.font = font_header
                        c.fill = fill_header
                        c.alignment = align_center
                        c.border = cell_border

                    start_row = 5
                    for row_idx in range(start_row, start_row + len(df_resumen)):
                        row_fill = fill_zebra if (row_idx % 2 == 0) else fill_white
                        for col_idx in range(1, 5):
                            c = ws.cell(row=row_idx, column=col_idx)
                            c.font = font_data
                            c.fill = row_fill
                            c.border = cell_border
                            if col_idx == 2: c.number_format = "@"
                            c.alignment = align_left if col_idx == 1 else align_center

                    for col in ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        for cell in col:
                            if cell.row in [1, 2]: continue
                            val_str = str(cell.value or "")
                            if len(val_str) > max_len: max_len = len(val_str)
                        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

            self.lbl_status_footer.configure(text="🟢 ¡Sincronización y Excel Creado!", text_color=COLORES["exito"])
            self.label_estado.configure(text="Todo guardado correctamente.", text_color=COLORES["exito"])
            self.card_container.configure(border_color=COLORES["exito"], border_width=2)
            
            if self.on_procesar_exito: self.on_procesar_exito()
            messagebox.showinfo("Éxito", "Los datos fueron procesados, guardados en la BD y el cuatrimestre se sincronizó automáticamente.")
        except Exception as e:
            messagebox.showerror("Error", f"Fallo al procesar:\n{str(e)}")