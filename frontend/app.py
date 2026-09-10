import os
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
import pandas as pd

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.modelos import guardar_ponchadas_desde_df


ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")


class AppLimpiador(ctk.CTk):

  def __init__(self):
    super().__init__()

    self.title("Procesador de Ponchadora con Formato Ejecutivo")
    self.geometry("600x470")
    self.resizable(False, False)

    self.ruta_archivo = None

    self.label_titulo = ctk.CTkLabel(
        self,
        text="Procesador de Ponchadora",
        font=ctk.CTkFont(size=22, weight="bold"),
    )
    self.label_titulo.pack(pady=(25, 5))

    self.label_subtitulo = ctk.CTkLabel(
        self,
        text="Corrección exacta de Hora local y Fechas México (DD/MM/AAAA).",
        font=ctk.CTkFont(size=12),
        text_color="gray",
    )
    self.label_subtitulo.pack(pady=(0, 20))

    self.frame_file = ctk.CTkFrame(self)
    self.frame_file.pack(pady=10, padx=30, fill="x")

    self.entry_ruta = ctk.CTkEntry(
        self.frame_file,
        placeholder_text="Ningún archivo seleccionado...",
        width=380,
    )
    self.entry_ruta.pack(side="left", padx=15, pady=15)

    self.btn_buscar = ctk.CTkButton(
        self.frame_file,
        text="Buscar CSV",
        width=100,
        command=self.seleccionar_archivo,
    )
    self.btn_buscar.pack(side="right", padx=15, pady=15)

    self.btn_procesar = ctk.CTkButton(
        self,
        text="Procesar y Generar Excel Ejecutivo",
        font=ctk.CTkFont(size=15, weight="bold"),
        fg_color="#1F4E78",
        hover_color="#153654",
        height=45,
        command=self.procesar_csv,
    )
    self.btn_procesar.pack(pady=25)

    self.label_estado = ctk.CTkLabel(
        self, text="", font=ctk.CTkFont(size=13, weight="bold")
    )
    self.label_estado.pack(pady=10)

  def seleccionar_archivo(self):
    archivo = filedialog.askopenfilename(
        title="Selecciona el CSV de la ponchadora",
        filetypes=[("Archivos CSV", "*.csv"), ("Todos los archivos", "*.*")],
    )
    if archivo:
      self.ruta_archivo = archivo
      self.entry_ruta.delete(0, tk.END)
      self.entry_ruta.insert(0, os.path.basename(archivo))
      self.label_estado.configure(
          text="Archivo listo para procesar.", text_color="white"
      )

  def procesar_csv(self):
    if not self.ruta_archivo:
      messagebox.showwarning(
          "Atención", "Por favor selecciona un archivo CSV primero."
      )
      return

    try:
      # 1. Cargar el CSV probando codificaciones
      try:
        df = pd.read_csv(
            self.ruta_archivo, sep=";", encoding="utf-8-sig", on_bad_lines="skip"
        )
      except Exception:
        df = pd.read_csv(
            self.ruta_archivo, sep=";", encoding="latin1", on_bad_lines="skip"
        )

      if df.shape[1] < 4:
        messagebox.showerror(
            "Error", "El archivo no tiene el formato esperado."
        )
        return

      # 2. Tomar Fecha_Hora (Columna 0) y Nombre (Columna 3)
      df_sub = df.iloc[:, [0, 3]].copy()
      df_sub.columns = ["Fecha_Hora_Raw", "Nombre"]

      df_sub["Nombre"] = df_sub["Nombre"].astype(str).str.strip()
      df_sub = df_sub.dropna(subset=["Fecha_Hora_Raw", "Nombre"])

      # 3. PARSEO EXACTO SIN CAMBIOS DE ZONA HORARIA
      df_sub["Fecha_Hora"] = pd.to_datetime(
          df_sub["Fecha_Hora_Raw"].astype(str).str.strip(),
          format="%d/%m/%Y %H:%M:%S",
          errors="coerce",
      )
      df_sub = df_sub.dropna(subset=["Fecha_Hora"])

      # Extraer Fecha en formato estricto México (DD/MM/AAAA)
      df_sub["Fecha_Texto"] = df_sub["Fecha_Hora"].dt.strftime("%d/%m/%Y")

      # 4. AGRUPACIÓN: Obtener entrada mínima y salida máxima reales
      df_resumen = (
          df_sub.groupby(["Fecha_Texto", "Nombre"], sort=False)
          .agg(
              Entrada=("Fecha_Hora", lambda x: x.min().strftime("%H:%M:%S")),
              Salida=("Fecha_Hora", lambda x: x.max().strftime("%H:%M:%S")),
          )
          .reset_index()
      )

      # Reordenar y renombrar columnas ANTES de registrar en la BD
      df_resumen.rename(columns={"Fecha_Texto": "Fecha"}, inplace=True)
      df_resumen = df_resumen[["Nombre", "Fecha", "Entrada", "Salida"]]

      # 5. PERSISTENCIA: Guardar automáticamente los datos extraídos en SQLite
      try:
        guardar_ponchadas_desde_df(df_resumen)
        print("Ponchadas guardadas con éxito en SQLite.")
      except Exception as e:
        print(f"Error al guardar en BD: {e}")

      # 6. Solicitar ubicación de guardado
      nombre_base = os.path.splitext(os.path.basename(self.ruta_archivo))[0]
      ruta_salida = filedialog.asksaveasfilename(
          title="Guardar Excel Formateado como...",
          initialfile=f"Reporte_Asistencia_{nombre_base}.xlsx",
          defaultextension=".xlsx",
          filetypes=[("Archivos de Excel", "*.xlsx")],
      )

      if not ruta_salida:
        return

      # 7. EXPORTAR A EXCEL Y APLICAR ESTILOS
      with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        df_resumen.to_excel(
            writer, sheet_name="Asistencia", index=False, startrow=3
        )
        ws = writer.sheets["Asistencia"]
        ws.views.sheetView[0].showGridLines = True

        # Estilos
        font_banner = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=11, color="1F2937")

        fill_banner = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        fill_header = PatternFill(
            start_color="2F5597", end_color="2F5597", fill_type="solid"
        )
        fill_zebra = PatternFill(
            start_color="F2F5F9", end_color="F2F5F9", fill_type="solid"
        )
        fill_white = PatternFill(
            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
        )

        align_center = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        align_left = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

        border_thin = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(
            left=border_thin,
            right=border_thin,
            top=border_thin,
            bottom=border_thin,
        )

        # BANNER DE ENCABEZADO
        ws.merge_cells("A1:D2")
        banner_cell = ws["A1"]
        banner_cell.value = "REPORTE CONSOLIDADO DE ASISTENCIA Y PONCHADAS"
        banner_cell.font = font_banner
        banner_cell.fill = fill_banner
        banner_cell.alignment = align_center

        # ENCABEZADOS DE LA TABLA
        for col_idx in range(1, 5):
          c = ws.cell(row=4, column=col_idx)
          c.font = font_header
          c.fill = fill_header
          c.alignment = align_center
          c.border = cell_border

        # DATOS
        start_row = 5
        total_rows = len(df_resumen)

        for row_idx in range(start_row, start_row + total_rows):
          row_fill = fill_zebra if (row_idx % 2 == 0) else fill_white

          for col_idx in range(1, 5):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.fill = row_fill
            c.border = cell_border

            # Asegurar formato de texto para fecha
            if col_idx == 2:
              c.number_format = "@"

            if col_idx == 1:
              c.alignment = align_left
            else:
              c.alignment = align_center

        # AUTO-AJUSTE ANCHO DE COLUMNAS
        for col in ws.columns:
          max_len = 0
          col_letter = get_column_letter(col[0].column)
          for cell in col:
            if cell.row in [1, 2]:
              continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
              max_len = len(val_str)
          ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

      self.label_estado.configure(
          text="¡Excel generado con éxito!", text_color="#2FA572"
      )
      messagebox.showinfo(
          "Éxito", f"Reporte formateado guardado en:\n{ruta_salida}"
      )

    except Exception as e:
      messagebox.showerror(
          "Error", f"Ocurrió un error al procesar el archivo:\n{str(e)}"
      )


if __name__ == "__main__":
  app = AppLimpiador()
  app.mainloop()